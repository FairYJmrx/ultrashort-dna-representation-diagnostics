from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def _find_project_root(start: Path) -> Path:
    for candidate in [start.parent, *start.parents]:
        if (candidate / "methods").is_dir() and (candidate / "configs").is_dir():
            return candidate
    raise RuntimeError("Could not locate the release repository root.")


PROJECT_ROOT = _find_project_root(Path(__file__).resolve())
MANUSCRIPT = PROJECT_ROOT / "manuscript"


INLINE_REPLACEMENTS = {
    r"\(x = x_1,\ldots,x_L\)": "x = x_1,...,x_L",
    r"\(x_i \in \{A,C,G,T,N\}\)": "x_i in {A,C,G,T,N}",
    r"\(w=x_i,\ldots,x_{i+k-1}\)": "w = x_i,...,x_{i+k-1}",
    r"\(c_w(x)\)": "c_w(x)",
    r"\(P=(p_1,\ldots,p_m)\)": "P = (p_1,...,p_m)",
    r"\(i\)": "i",
    r"\(c_{\operatorname{canon}(s_{i,P})}(x)\)": "c_{canon(s_i,P)}(x)",
    r"\(P=(0,2,4,6)\)": "P = (0,2,4,6)",
    r"\(g(x)\)": "g(x)",
}


def clear_paragraph_content(paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)
    paragraph._p = paragraph._element = None


def math_run(text: str):
    run = OxmlElement("m:r")
    text_node = OxmlElement("m:t")
    text_node.text = text
    run.append(text_node)
    return run


def math_subscript(base: str, sub: str):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    e.append(math_run(base))
    s = OxmlElement("m:sub")
    s.append(math_run(sub))
    node.append(e)
    node.append(s)
    return node


def set_math_paragraph(paragraph, pieces) -> None:
    clear_paragraph_content(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)
    math_para = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    for piece in pieces:
        if isinstance(piece, tuple) and piece[0] == "sub":
            math.append(math_subscript(piece[1], piece[2]))
        else:
            math.append(math_run(str(piece)))
    math_para.append(math)
    paragraph._p.append(math_para)


def replace_inline_latex(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        text = paragraph.text
        new = text
        for old, repl in INLINE_REPLACEMENTS.items():
            new = new.replace(old, repl)
        if new != text:
            clear_paragraph_content(paragraph)
            run = paragraph.add_run(new)
            run.font.name = "Times New Roman"
            run.font.size = Pt(10.5)


def replace_display_equations(doc: Document) -> None:
    blocks = []
    paragraphs = list(doc.paragraphs)
    idx = 0
    while idx < len(paragraphs) - 2:
        if paragraphs[idx].text.strip() == r"\[" and paragraphs[idx + 2].text.strip() == r"\]":
            blocks.append((paragraphs[idx], paragraphs[idx + 1], paragraphs[idx + 2]))
            idx += 3
        else:
            idx += 1

    equations = [
        ["canon(w) = ", ("sub", "min", "lex"), "(w, rc(w))."],
        [("sub", "s", "i,P"), "(x) = ", ("sub", "x", "i+p1"), " ... ", ("sub", "x", "i+pm"), "."],
        [("sub", "蠁", "CSP"), "(x) = L2([L2(", ("sub", "c", "canon-spaced"), "(x)); g(x)])."],
    ]
    for block, equation in zip(blocks, equations):
        open_para, equation_para, close_para = block
        set_math_paragraph(open_para, equation)
        delete_paragraph(equation_para)
        delete_paragraph(close_para)


def patch_docx(path: Path) -> None:
    doc = Document(str(path))
    replace_inline_latex(doc)
    replace_display_equations(doc)
    doc.save(str(path))


def main() -> None:
    stage3 = MANUSCRIPT / "stage3_manuscript_v4.docx"
    final = MANUSCRIPT / "final_manuscript.docx"
    patch_docx(stage3)
    shutil.copy2(stage3, final)
    shutil.copy2(MANUSCRIPT / "stage3_manuscript_v4.md", MANUSCRIPT / "final_manuscript.md")
    print(f"Patched equations in {stage3}")
    print(f"Copied patched manuscript to {final}")


if __name__ == "__main__":
    main()

