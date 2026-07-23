from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def _find_project_root(start: Path) -> Path:
    for candidate in [start.parent, *start.parents]:
        if (candidate / "methods").is_dir() and (candidate / "configs").is_dir():
            return candidate
    raise RuntimeError("Could not locate the release repository root.")


PROJECT_ROOT = _find_project_root(Path(__file__).resolve())
MANUSCRIPT = PROJECT_ROOT / "manuscript"
TABLES = MANUSCRIPT / "tables"
FIGS = MANUSCRIPT / "figures"
REFERENCES = PROJECT_ROOT / "references" / "references.bib"

TITLE = "Controlled information-preservation diagnostics for ultra-short DNA read representations"
SUBTITLE = (
    "Canonical k-mers preserve high-resolution identity evidence, whereas canonical spaced-property encoding "
    "provides compact perturbation-stable auxiliary evidence for mNGS-like reads"
)
AUTHORS = "MEI Ruixiang"
REPO_URL = "https://github.com/FairYJmrx/ultrashort-dna-representation-diagnostics"
SUBMISSION_TAG = "v0.2.1-stage2-mei-submission"
MD_PATH = MANUSCRIPT / "stage2_manuscript_v2.md"
DOCX_PATH = MANUSCRIPT / "stage2_manuscript_v2.docx"


def read_markdown_table(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    lines = [line for line in path.read_text(encoding="utf-8").splitlines() if line.strip().startswith("|")]
    if len(lines) < 3:
        return pd.DataFrame()
    header = [cell.strip() for cell in lines[0].strip("|").split("|")]
    rows = []
    for line in lines[2:]:
        rows.append([cell.strip() for cell in line.strip("|").split("|")])
    return pd.DataFrame(rows, columns=header)


def df_to_markdown(df: pd.DataFrame) -> str:
    if df.empty:
        return ""
    return df.to_markdown(index=False)


def parse_markdown_table_lines(lines: list[str]) -> pd.DataFrame:
    if len(lines) < 3:
        return pd.DataFrame()
    header = [cell.strip() for cell in lines[0].strip("|").split("|")]
    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in lines[2:]]
    return pd.DataFrame(rows, columns=header)


def select_rows(df: pd.DataFrame, n: int) -> pd.DataFrame:
    return df.head(n).copy() if not df.empty else df


def bib_entries() -> dict[str, dict[str, str]]:
    text = REFERENCES.read_text(encoding="utf-8")
    blocks = re.split(r"\n(?=@)", text.strip())
    entries: dict[str, dict[str, str]] = {}
    for block in blocks:
        key_match = re.match(r"@\w+\{([^,]+),", block)
        if not key_match:
            continue
        key = key_match.group(1)
        entry: dict[str, str] = {"key": key}
        for field in ["author", "title", "journal", "year", "volume", "pages", "doi", "eprint", "archivePrefix"]:
            match = re.search(rf"\n\s*{field}\s*=\s*\{{(.*?)\}}\s*,?", block, re.IGNORECASE | re.DOTALL)
            if match:
                entry[field] = re.sub(r"\s+", " ", match.group(1)).strip()
        entries[key] = entry
    return entries


def author_year(entry: dict[str, str]) -> str:
    authors = entry.get("author", entry.get("key", "")).split(" and ")
    first = authors[0].split(",")[0]
    if len(authors) == 1:
        author_part = first
    elif len(authors) == 2:
        author_part = f"{first} and {authors[1].split(',')[0]}"
    else:
        author_part = f"{first} et al."
    return f"{author_part}, {entry.get('year', 'n.d.')}"


def format_reference(entry: dict[str, str]) -> str:
    authors = entry.get("author", "").split(" and ") if entry.get("author") else []
    if len(authors) > 8:
        author_text = "; ".join(authors[:8]) + "; et al."
    else:
        author_text = "; ".join(authors)
    parts = []
    if author_text:
        parts.append(author_text)
    if entry.get("year"):
        parts.append(f"({entry['year']}).")
    if entry.get("title"):
        parts.append(f"{entry['title']}.")
    if entry.get("journal"):
        parts.append(f"{entry['journal']}.")
    if entry.get("volume"):
        parts.append(f"{entry['volume']}.")
    if entry.get("pages"):
        parts.append(f"{entry['pages']}.")
    if entry.get("doi"):
        parts.append(f"https://doi.org/{entry['doi']}")
    elif entry.get("eprint"):
        parts.append(f"{entry.get('archivePrefix', 'arXiv')}:{entry['eprint']}")
    return " ".join(part for part in parts if part).strip()


class CitationManager:
    def __init__(self) -> None:
        self.entries = bib_entries()
        self.used: list[str] = []

    def cite(self, *keys: str) -> str:
        labels = []
        for key in keys:
            if key not in self.entries:
                raise KeyError(f"Missing reference key: {key}")
            if key not in self.used:
                self.used.append(key)
            labels.append(author_year(self.entries[key]))
        return "(" + "; ".join(labels) + ")"

    def references_markdown(self) -> str:
        lines = []
        for idx, key in enumerate(self.used, start=1):
            lines.append(f"{idx}. {format_reference(self.entries[key])}")
        return "\n\n".join(lines)


def make_formula_block() -> str:
    return r"""
Let a DNA read be \(x = x_1,\ldots,x_L\), with \(x_i \in \{A,C,G,T,N\}\). For a contiguous word \(w=x_i,\ldots,x_{i+k-1}\), ordinary k-mer counting stores \(c_w(x)\). The reverse-complement canonical operation maps a word and its reverse complement to one feature index,

\[
\operatorname{canon}(w)=\min_{\mathrm{lex}}(w,\operatorname{rc}(w)).
\]

For a spaced seed pattern \(P=(p_1,\ldots,p_m)\), the spaced token beginning at position \(i\) is

\[
s_{i,P}(x)=x_{i+p_1}\ldots x_{i+p_m}.
\]

The canonical spaced count vector is \(c_{\operatorname{canon}(s_{i,P})}(x)\). CSP uses the default pattern \(P=(0,2,4,6)\), then concatenates a low-dimensional property vector \(g(x)\). The property vector contains the mean and standard deviation of hydrogen-bond class, GC indicator, purine indicator and EIIP-like base values, plus N fraction, normalized length and Shannon entropy. The final representation is

\[
\phi_{\mathrm{CSP}}(x)=\operatorname{L2}\left([\operatorname{L2}(c_{\operatorname{canon-spaced}}(x)); g(x)]\right).
\]

Thus CSP is not a learned embedding. It is a deterministic auxiliary block that couples strand-canonical spaced evidence with interpretable biochemical summaries.
"""


def build_markdown() -> str:
    c = CitationManager()

    hospital = read_markdown_table(TABLES / "stage2_table_hospital_69_75_focus.md")
    gain = read_markdown_table(TABLES / "stage2_table_stability_gain_summary.md")
    csp_full = read_markdown_table(TABLES / "stage2_table_csp_full_ablation.md")
    csp_components = read_markdown_table(TABLES / "stage2_table_csp_component_singletons.md")
    readout_agg = read_markdown_table(TABLES / "stage2_table_readout_aggregate.md")
    attention_cp = read_markdown_table(TABLES / "stage2_table_attention_change_points.md")
    param_best = read_markdown_table(TABLES / "stage2_table_parameter_readout_best.md")
    neural_agg = read_markdown_table(TABLES / "stage2_table_neural_compatibility_aggregate.md")
    arg_readout = read_markdown_table(TABLES / "stage2_table_arg_snp_readout_aggregate.md")
    arg_stability = read_markdown_table(TABLES / "stage2_table_arg_snp_best_stability.md")

    csp_full_focus = csp_full[
        csp_full["Condition"].isin(["3% N mask", "6-bp local mismatch", "1% substitution + 3% N"])
        & csp_full["Length"].astype(str).isin(["69", "75", "150"])
    ].copy()
    arg_stability_focus = arg_stability[
        arg_stability["Boundary task"].isin(["arg_family", "arg_allele", "resistance_snp"])
        & arg_stability["Condition"].isin(["3% N mask", "6-bp local mismatch"])
        & arg_stability["Length"].astype(str).isin(["69", "75"])
    ].copy()

    md: list[str] = []
    md.append(f"# {TITLE}\n")
    md.append(f"**{SUBTITLE}**\n")
    md.append(f"Author: {AUTHORS}\n")

    md.append("## Abstract\n")
    md.append(
        "Clinical metagenomic next-generation sequencing (mNGS) often produces short or quality-trimmed reads, yet "
        "DNA representations are still commonly judged by downstream accuracy alone. This can obscure which information "
        "an encoding preserves before any classifier is trained. We present a controlled representation-diagnostics "
        "framework for ultra-short DNA reads and define canonical spaced-property encoding (CSP), a deterministic feature "
        "block that combines reverse-complement canonical spaced-seed counts with interpretable biochemical summaries. "
        "Across a close-relative WGS-slice grid spanning 69, 75, 100, 110, 125, 150 bp and a PE150 proxy, CSP was the top "
        "clean-perturbed stability representation in 42 of 42 length-by-perturbation settings. At 69 bp, CSP achieved "
        "mean paired cosine of 0.994 under 3% N masking and 0.988 under a 6-bp local mismatch, compared with 0.940 and "
        "0.877 for canonical 5-mers. This advantage did not translate into universal species or resistance accuracy: "
        "canonical k-mers remained strong high-resolution identity baselines in close-relative and ARG/SNP probes. A "
        "dense context-visibility diagnostic further showed that apparent read-length thresholds depend on motif position, "
        "with motif-pair visibility emerging at 130, 140, 148 or 155 bp under different placements. Deterministic neural "
        "probes further showed that small CNN and tiny Transformer readouts were task-dependent rather than universally "
        "superior. These results support a bounded conclusion: short-read DNA pipelines should separate exact identity "
        "evidence from compact perturbation-stable auxiliary evidence, rather than ranking representations by a single "
        "accuracy number."
    )

    md.append("## Introduction\n")
    md.append(
        f"Clinical mNGS has become an important route for pathogen detection because it can detect unexpected organisms "
        f"without a fixed target panel {c.cite('Wilson2014', 'Wilson2019', 'Chiu2019')}. The same clinical setting also "
        f"creates difficult input conditions for computational analysis. Reads may be shortened by adapter and quality "
        f"trimming, contain ambiguous bases, appear from either strand, or include background and low-biomass artifacts "
        f"{c.cite('Martin2011', 'Bolger2014', 'Salter2014')}. These constraints make the representation layer more than "
        f"an engineering detail. Before a classifier, aligner or database index can succeed, the encoding has already "
        f"decided which sequence properties remain available."
    )
    md.append(
        f"Most mature metagenomic classifiers are built around exact or near-exact word evidence. Kraken, Kraken 2, "
        f"CLARK, Centrifuge and Kaiju show the practical power of indexed k-mer, minimizer or translated-word matching "
        f"at scale {c.cite('Wood2014', 'Wood2019', 'Ounit2015', 'Kim2016', 'Menzel2016')}. CAMI benchmarks further show "
        f"that apparent performance depends on novelty, taxonomic difficulty, abundance structure and database coverage "
        f"{c.cite('Sczyrba2017', 'Meyer2022')}. These observations argue against using a small local study as a clinical "
        f"species-identification leaderboard. They also motivate a narrower and testable question: what information does "
        f"each representation preserve under the short-read perturbations that clinical pipelines actually encounter?"
    )
    md.append(
        f"Controlled simulated reads are used here as a methodological choice, not as a substitute claim for clinical "
        f"diagnostic validation. Real clinical mNGS datasets often provide sample-level diagnostic or qPCR information, "
        f"but they rarely provide read-level truth, exact sequencing-error origin, mutation status and controlled read length "
        f"for every fragment. WGS-derived reads therefore support mechanism-level decomposition because source genome, read "
        f"length, perturbation and clean-perturbed pairing are known. ART Illumina simulation adds a field-standard sequencing "
        f"error model with platform-like quality profiles {c.cite('Huang2012ART')}, and CAMI low-complexity data provide an "
        f"external metagenomic benchmark context {c.cite('Sczyrba2017', 'Meyer2022')}. This three-layer design is intended to "
        f"close the representation-diagnostics logic while reserving end-to-end clinical mNGS pipelines for future work."
    )
    md.append(
        f"Deep learning has widened the representational vocabulary for DNA. Convolutional and recurrent models have "
        f"been used to learn regulatory sequence specificity {c.cite('Alipanahi2015', 'Zhou2015', 'Quang2016')}; "
        f"read-level metagenomic neural classifiers include recurrent and attention-based models "
        f"{c.cite('Liang2020', 'Wichmann2023')}; and DNA foundation models now include k-mer token models, efficient "
        f"multi-species pretraining, long-context models, reverse-complement-aware architectures and single-nucleotide "
        f"generative models {c.cite('Ji2021', 'Zhou2024', 'DallaTorre2025', 'Nguyen2023', 'Schiff2024', 'Fishman2025', 'Nguyen2024Evo')}. "
        f"However, a more expressive model cannot attend to information that is absent from the observed read. The "
        f"short-read setting therefore requires diagnostics that separate model capacity from input information loss."
    )
    md.append(
        "Here we propose a controlled information-preservation study rather than a final clinical classifier. Our central "
        "claim is not that CSP replaces canonical k-mers. Instead, canonical k-mers provide high-resolution identity "
        "evidence, whereas CSP provides compact, strand-friendly and perturbation-stable auxiliary evidence. We evaluate "
        "this claim using WGS-derived close-relative reads, hospital-like 69/75 bp perturbation grids, CSP component "
        "ablation, k and spaced-pattern sensitivity, lightweight readout probes, deterministic neural compatibility probes, "
        "synthetic ARG/SNP boundary tasks and attention-style context-visibility diagnostics."
    )

    md.append("## Related Work\n")
    md.append(
        f"Alignment-free sequence comparison treats word content as a proxy for sequence relatedness. k-mer counting, "
        f"MinHash sketches and related methods provide efficient representations for genome comparison and metagenomic "
        f"classification {c.cite('Marcais2011', 'Ondov2016', 'Zielezinski2017')}. Canonical k-mers collapse a word and "
        f"its reverse complement to the same index, which is useful for strand-ambiguous reads but can remove strand-specific "
        f"signals. Spaced seeds, introduced for sensitive homology search and later applied to metagenomic classification, "
        f"sample non-contiguous positions within a word and can improve tolerance to mismatches {c.cite('Ma2002', 'Brinda2015')}."
    )
    md.append(
        f"DNA can also be represented as a numerical signal. Shannon information theory provides a language for uncertainty "
        f"and information loss {c.cite('Shannon1948')}, while chaos game representation, genomic signal processing and "
        f"EIIP-style mappings show that nucleotide sequences can be converted into compositional or physicochemical channels "
        f"{c.cite('Jeffrey1990', 'Voss1992', 'Anastassiou2001', 'Cristea2002', 'Nair2006')}. dna2vec similarly bridges "
        f"discrete k-mers and continuous representations {c.cite('Ng2017')}. CSP follows this tradition in a deliberately "
        f"modest way: it adds interpretable biochemical summaries to canonical spaced counts instead of learning a new "
        f"embedding from large corpora."
    )
    md.append(
        f"Transformer models add the separate issue of token position and co-occurrence. Self-attention can connect all "
        f"observed tokens {c.cite('Vaswani2017')}, and rotary position embeddings provide a compact relative-position "
        f"mechanism {c.cite('Su2021')}. But attention cannot recover a motif that is outside the sequenced fragment. We "
        f"therefore include an attention-context diagnostic that measures visibility of motif relations before attributing "
        f"read-length effects to a particular neural architecture."
    )
    md.append(
        f"ARG and antimicrobial-resistance analysis imposes stricter biological requirements than coarse taxonomic assignment. "
        f"Resources and tools such as CARD, AMRFinderPlus, ResFinder and MEGARes/AMR++ encode curated gene families, "
        f"protein-level evidence, mutation rules and resistome workflows {c.cite('Alcock2023', 'Feldgarden2021', 'Bortolaia2020', 'Bonin2023')}. "
        f"DeepARG illustrates the use of learned models for ARG prediction {c.cite('ArangoArgoty2018')}. Our experiments "
        f"do not claim clinical ARG calling. They test whether a compact property-aware block can preserve perturbed "
        f"ARG-like signal and where exact sequence evidence remains indispensable."
    )

    md.append("## Terminology and Contribution\n")
    md.append(
        "We use one term consistently throughout the manuscript. **Canonical k-mer** denotes reverse-complement pooled "
        "contiguous k-mer counts. **Canonical spaced seed** denotes reverse-complement pooled counts of non-contiguous "
        "tokens. **CSP** denotes canonical spaced-property encoding, implemented in code as `cspaced_property_l2`. "
        "**Hybrid** denotes a concatenation of canonical k-mer counts and CSP followed by L2 normalization. **Perturbation "
        "stability** denotes similarity between a clean read and a mutated, N-masked, trimmed or locally mismatched version "
        "of the same read. It does not mean resistance to multi-species contamination."
    )
    md.append(
        "The contribution is therefore a representation-diagnostics framework, not a new end-to-end taxonomic classifier. "
        "The framework is designed to answer four questions: (i) which features remain stable when a short read is lightly "
        "perturbed, (ii) which priors inside CSP contribute to stability, (iii) whether identity-like readout probes favor "
        "the same representation, and (iv) when short reads structurally remove context that an attention model would need."
    )

    md.append("## Method\n")
    md.append("### Rationale for controlled simulated reads\n")
    md.append(
        "The data design follows the level of the claim. Because this manuscript tests representation-level information "
        "preservation rather than end-to-end clinical classification, the primary experiments require read-level ground "
        "truth and paired clean-versus-perturbed fragments. WGS-derived reads provide controlled species origin, read length "
        "and perturbation axes. ART Illumina validation is planned to test whether the same stability patterns hold under "
        "a field-standard sequencing-error profile rather than only under hand-specified substitutions, N masking and trimming. "
        "CAMI low-complexity data are planned as an external benchmark probe for lightweight readout, not as a production "
        "taxonomic-classification leaderboard."
    )
    md.append("### Representation definitions\n")
    md.append(make_formula_block())
    md.append(
        "Canonical 5-mer and 7-mer baselines were included as high-resolution identity features. Canonical spaced seeds "
        "were included as a compact mismatch-tolerant baseline. Hybrid features were constructed as L2-normalized "
        "concatenations of canonical k-mer counts and CSP. All vocabulary-dependent features were fitted on the training "
        "or clean subset defined by each experiment to avoid using perturbed test sequences to define the vocabulary."
    )
    md.append("### Data sources and perturbations\n")
    md.append(
        "The main stage-2 WGS panel contained 21 genomes from six close or clinically relevant genera: Acinetobacter, "
        "Burkholderia, Candida, Enterobacter, Escherichia and Klebsiella. Reads were generated at 69, 75, 100, 110, "
        "125 and 150 bp, plus a PE150 proxy represented as a 300 bp paired-end-equivalent window. Each length contained "
        "1,680 clean reads before perturbation. Perturbations were generated with fixed random seeds and included 1% "
        "substitution, 3% N masking, 5-bp trimming, combined 1% substitution plus 3% N masking, short indels and a 6-bp "
        "local mismatch block."
    )
    md.append(
        "The 69/75 bp analysis was treated as a hospital-like short-read setting because the user-facing project context "
        "emphasized 75 bp single-end reads and approximately 69 bp post-QC reads. This experiment measured whether a "
        "perturbed read stayed close to its clean counterpart, not whether a clinical sample with multiple organisms was "
        "classified correctly."
    )
    md.append(
        "Stage-3 validation extends this data hierarchy without changing the paper's scope. ART Illumina profiles will be "
        "used to generate platform-like sequencing-error reads from the same reference genomes, preserving the stability "
        "metrics while replacing hand-specified perturbations with a commonly used read simulator. CAMI low-complexity "
        "reads will be processed as an external benchmark subset for target/background or genus-level lightweight readout "
        "probes. These additions are designed to test externality and noise-model robustness, not to claim clinical sensitivity "
        "or specificity."
    )
    md.append("### Metrics and readout probes\n")
    md.append(
        "Perturbation stability was measured by paired clean-perturbed cosine similarity, paired L2 drift and nearest-clean "
        "retrieval. Compactness was measured by feature dimension and density. Readout probes used nearest centroid, "
        "logistic regression and a small scikit-learn MLP with fixed random seeds. These probes measured whether a signal "
        "could be extracted by simple models. They were not interpreted as clinical accuracy estimates."
    )
    md.append(
        "CSP ablation separated canonical spaced counts from property additions: hydrogen-bond class, GC indicator, purine "
        "indicator, EIIP-like values, N fraction, entropy and length. Parameter sensitivity swept canonical k-mer values "
        "from k=4 to k=9 and several spaced seed patterns. The attention-context diagnostic varied read length densely "
        "from 110 to 160 bp and moved a class-defining motif pair across positions 120, 130, 138 and 145. Synthetic ARG/SNP "
        "boundary probes tested ARG-family, ARG-allele and resistance-SNP style tasks under the same perturbation logic."
    )
    md.append(
        "The neural compatibility probe used PyTorch CPU with deterministic seeds and single-thread execution. It compared "
        "tabular MLP readouts for canonical 5-mer, CSP and hybrid vectors with 1D-CNN and one-layer tiny Transformer readouts "
        "over one-hot or property channels. The probe covered 69, 75, 100 and 150 bp reads, clean/N-masked/combined-perturbation "
        "conditions, target/background classification, global species stress classification and an Enterobacter within-genus "
        "species task. It was designed to test model-readability, not clinical accuracy."
    )

    md.append("## Results\n")
    md.append("### CSP had its clearest advantage in clean-perturbed stability\n")
    md.append(
        "Across the WGS-slice perturbation grid, CSP was the top representation by paired clean-perturbed cosine in all "
        "42 length-by-perturbation settings. The advantage was largest in short and locally disrupted reads. At 69 bp "
        "with a 6-bp local mismatch, CSP reached mean paired cosine of 0.988 and mean L2 drift of 0.157, whereas canonical "
        "5-mer reached 0.877 and 0.494. Under 3% N masking at 69 bp, CSP reached 0.994 paired cosine and 0.107 L2 drift, "
        "whereas canonical 5-mer reached 0.940 and 0.344. CSP also had far fewer features than canonical 7-mers, which "
        "used roughly 8,000 observed features in the 69/75 bp WGS grid."
    )
    md.append(df_to_markdown(select_rows(hospital, 20)))
    md.append("\n![Figure 1. Clean-perturbed feature stability across read length.](figures/stage2_fig_stability_grid.png)\n")
    md.append("\n![Figure 2. Hospital-like 69/75 bp perturbation drift.](figures/stage2_fig_hospital_69_75_l2.png)\n")
    md.append(
        "The summary comparison confirmed that this was not a single-condition artifact. CSP exceeded canonical 5-mer, "
        "canonical 7-mer, canonical spaced seed and canonical 5-mer+CSP in paired-cosine stability in all matched comparisons."
    )
    md.append(df_to_markdown(gain))

    md.append("### Ablation showed that the property block, especially hydrogen-bond and entropy summaries, contributed to stability\n")
    md.append(
        "CSP's stability advantage was not attributable to one isolated scalar. The full property block improved stability "
        "over canonical spaced seed counts across the tested short-read perturbations. In singleton ablations, hydrogen-bond "
        "class had the largest mean cosine gain over canonical spaced seeds, followed by entropy, purine and GC summaries. "
        "N fraction alone added little in this setup, which is expected because N masking also changes the count space."
    )
    md.append(df_to_markdown(csp_components))
    md.append(df_to_markdown(select_rows(csp_full_focus, 12)))
    md.append("\n![Figure 3. Singleton property ablation.](figures/stage2_fig_csp_singleton_ablation.png)\n")

    md.append("### Lightweight readout probes separated robustness from identity resolution\n")
    md.append(
        "Readout probes did not reproduce the stability ranking as a universal accuracy ranking. In the within-genus "
        "species probe, canonical 5-mer had the highest mean macro-F1 among the tested stage-2 representations, followed "
        "closely by canonical 7-mer and the canonical 5-mer+CSP hybrid. CSP was lower. In the target/background probe, "
        "CSP had the highest mean macro-F1 but the absolute scores remained modest. This distinction is central: CSP "
        "preserved perturbed information well, but exact k-mer evidence remained important for fine identity resolution."
    )
    md.append(df_to_markdown(readout_agg))
    md.append("\n![Figure 4. Lightweight readout probes remained task-dependent.](figures/stage2_fig_readout_aggregate.png)\n")

    md.append("### k and spaced-pattern sensitivity argued against a single-parameter recommendation\n")
    md.append(
        "The parameter grid showed that the best readout configuration changed with task and read length. Target/background "
        "at 69 bp favored a canonical spaced pattern, target/background at 75 bp favored k=5, and within-genus species at "
        "150 bp favored canonical k=6. This supports the paper's framing as a representation-diagnostics study rather than "
        "a universal prescription for one k or one seed pattern."
    )
    md.append(df_to_markdown(param_best))

    md.append("### Deterministic neural probes showed model compatibility, not neural superiority\n")
    md.append(
        "The additional PyTorch probe trained 252 small neural readouts with fixed seeds. It did not support a broad claim "
        "that CNNs or tiny Transformers automatically improve ultra-short read interpretation. In target/background probes, "
        "1D-CNN over one-hot channels had the highest mean macro-F1, while CSP read by a tabular MLP was close and used only "
        "147 features on average. In global species and within-genus Enterobacter stress probes, the best small models were "
        "tabular MLPs over canonical or hybrid vectors, and absolute macro-F1 values remained low. This supports a practical "
        "model-matching interpretation: CSP is a natural compact tabular auxiliary input, whereas one-hot or property channels "
        "are more appropriate when a CNN or attention model is explicitly trained."
    )
    md.append(df_to_markdown(neural_agg))
    md.append("\n![Figure 5. Deterministic neural compatibility probes.](figures/stage2_fig_neural_compatibility.png)\n")

    md.append("### Context loss around 125-150 bp was position-dependent, not a single read-length threshold\n")
    md.append(
        "The attention-context diagnostic directly addressed the concern that a coarse 125 versus 150 bp comparison could "
        "misidentify a threshold. The first length with full motif-pair visibility shifted with motif placement: 130 bp "
        "when the motif pair began near position 120, 140 bp near position 130, 148 bp near position 138 and 155 bp near "
        "position 145. Thus the loss of context is not simply proportional to the number of missing bases. If a biologically "
        "or semantically relevant motif relation lies outside the read, self-attention can still connect observed tokens "
        "but cannot model the missing relation."
    )
    md.append(df_to_markdown(attention_cp))
    md.append("\n![Figure 6. Attention-style context visibility breakpoints.](figures/stage2_fig_attention_breakpoints.png)\n")
    md.append("\n![Figure 7. Best readout transitions around motif visibility.](figures/stage2_fig_attention_f1_breakpoints.png)\n")

    md.append("### ARG/SNP boundary probes bounded the role of CSP\n")
    md.append(
        "Synthetic ARG/SNP probes showed why CSP should be treated as auxiliary evidence. CSP was again frequently the top "
        "stability representation, but the readout tasks were too easy for many representations in ARG-family and ARG-allele "
        "settings. The resistance-SNP probe showed closer differences among methods. These results do not establish clinical "
        "ARG or SNP calling. They support a weaker but useful interpretation: CSP can preserve perturbed ARG-like feature "
        "proximity, while exact k-mer, alignment or curated database evidence remains necessary for allele-level and SNP-level "
        "decisions."
    )
    md.append(df_to_markdown(select_rows(arg_stability_focus, 18)))
    md.append(df_to_markdown(arg_readout))
    md.append("\n![Figure 8. ARG/SNP boundary readout probes.](figures/stage2_fig_arg_snp_readout.png)\n")

    md.append("### Planned stage-3 external validation will test simulator and benchmark generality\n")
    md.append(
        "[Planned validation: ART Illumina error-profile results will be inserted here after the stage-3 run. This analysis "
        "will compare canonical k-mer, canonical spaced seed, CSP, hybrid, MinHash sketch and EIIP baselines using paired "
        "cosine, L2 drift and nearest-clean retrieval under ART-generated Illumina-like sequencing errors.]"
    )
    md.append(
        "[Planned validation: compact classical baseline results will be inserted here after the stage-3 run. MinHash sketches "
        "will test whether a compact alignment-free sketch explains the CSP advantage, and EIIP-only baselines will test "
        "whether stability comes merely from using a numerical DNA signal.]"
    )
    md.append(
        "[Planned validation: CAMI low-complexity readout-probe results will be inserted here after the stage-3 run. The CAMI "
        "analysis will be treated as an external benchmark probe for signal readability and degradation, not as an end-to-end "
        "clinical or SOTA metagenomic classifier comparison.]"
    )

    md.append("## Discussion\n")
    md.append(
        "The main result is a division of labor among representations. Canonical k-mers are still the most defensible "
        "backbone for exact identity evidence, especially when the task is close species, strain, allele or SNP resolution. "
        "CSP contributes a different property: it keeps perturbed short reads close to their clean counterparts in a compact "
        "and interpretable space. In practical terms, this means CSP is better framed as an auxiliary robustness block, "
        "a QC/audit feature or a dense side channel for downstream models, not as a replacement for canonical k-mer indices."
    )
    md.append(
        "This distinction also resolves the apparent conflict around accuracy. Accuracy and macro-F1 are useful only as "
        "readout probes in this paper. They ask whether a simple model can extract a signal from the representation. They "
        "do not estimate clinical sensitivity, specificity or diagnostic accuracy. Overemphasizing accuracy would be "
        "misleading because the panel is intentionally controlled and small. Underemphasizing all readout probes would also "
        "be incomplete because a representation that preserves information but cannot be read by any downstream model would "
        "have limited practical value."
    )
    md.append(
        "The most realistic future route is hybrid evidence. For mNGS species identification, canonical k-mers, alignment "
        "or database indices should provide high-resolution taxonomic evidence, whereas CSP can track whether short, N-masked "
        "or locally mismatched reads remain compositionally and biochemically near the expected clean signal. For ARG work, "
        "CSP may help characterize degraded or ambiguous reads and expose interpretable shifts, but allele calling, resistance "
        "SNP interpretation, gene context and plasmid linkage require exact sequence, protein-domain or curated database evidence."
    )
    md.append(
        "Several conclusions remain deliberately unproven. The neural probe was intentionally small and local; it does not "
        "establish CNN or Transformer superiority on realistic clinical mNGS data. The stage-3 ART and CAMI analyses are "
        "planned to test measurement robustness and external benchmark readability, but they still will not constitute "
        "clinical sensitivity, specificity or production-pipeline validation. Kraken2, Centrifuge, Kaiju, alignment pipelines "
        "and curated CARD/ResFinder/AMRFinderPlus tasks remain the appropriate next layer once the representation-level "
        "evidence is fixed."
    )

    md.append("## Future Work\n")
    md.append(
        "Future work will extend the proposed layered evidence architecture into an end-to-end metagenomic workflow. In that "
        "larger system, canonical k-mer, alignment or database evidence should provide high-resolution taxonomic or ARG identity "
        "support, while CSP-like auxiliary features provide compact perturbation-stability, quality-audit and confidence-side "
        "information for degraded, N-masked, trimmed or locally mismatched reads. This next stage should use realistic FASTQ "
        "quality profiles, host/background mixtures, abundance variation, larger genome-held-out panels and eventually real "
        "clinical mNGS samples with sample-level orthogonal validation."
    )

    md.append("## Limitations\n")
    md.append(
        "The WGS panel contained 21 genomes from six genera and was not designed to represent microbial diversity, hospital "
        "background mixtures, abundance variation, host depletion, database incompleteness, sample-level uncertainty or wet-lab "
        "contamination. The hand-specified perturbations model substitutions, N masking, trimming, short indels and local "
        "mismatches for mechanism decomposition, but they do not reproduce a complete sequencing run or library-preparation "
        "process. ART Illumina validation is planned to add a field-standard error-profile layer, yet even ART cannot replace "
        "real clinical host background, abundance structure, contamination and diagnostic uncertainty. CAMI low-complexity "
        "analysis is planned as an external benchmark probe, not a clinical endpoint. The readout models are intentionally "
        "small and deterministic. Therefore, CSP-alone species identification, ARG allele calling, resistance SNP classification, "
        "mobile-element context and plasmid linkage should not be claimed from these data."
    )

    md.append("## Code and Data Availability\n")
    md.append(
        "All code, generated lightweight reads, result tables, figures and manuscript builders are maintained in the project "
        f"repository ({REPO_URL}; release tag {SUBMISSION_TAG}; exact commit hash to be reported from the public archive or "
        "cover letter at submission). Random seeds are fixed in the stage-2 scripts, "
        "and the earlier manuscript/results snapshot was preserved as an internal project archive before the stage-2 rerun. "
        "The submitted code package includes executable scripts, configuration files, generated summary tables, figures, "
        "manuscript builders and the 21-genome close-relative WGS-slice manifest; bulky downloaded reference FASTA files "
        "are intentionally excluded and can be regenerated from the manifest and preparation scripts. If journal policy "
        "requires public access, the private repository should be made public or archived with a DOI after double-blind "
        "constraints are resolved."
    )

    md.append("## Conclusions\n")
    md.append(
        "No single DNA representation dominated all short-read mNGS-like settings. Canonical k-mers remained the strongest "
        "general-purpose identity evidence. CSP provided a compact, strand-friendly and perturbation-stable auxiliary "
        "representation, with its strongest evidence in 69/75 bp N masking, local mismatch and combined perturbation settings. "
        "The paper's actionable message is therefore architectural rather than competitive: short-read pipelines should layer "
        "exact identity evidence with auxiliary stability evidence, and should evaluate read length through explicit context-"
        "visibility diagnostics when attention-like models are considered."
    )

    md.append("## References\n")
    md.append(c.references_markdown())
    return "\n\n".join(md).strip() + "\n"


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) < 16 else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(8)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table_to_doc(doc: Document, df: pd.DataFrame, max_rows: int = 18) -> None:
    if df.empty:
        return
    if len(df.columns) > 8 or len(df) > 18:
        note = doc.add_paragraph(
            f"Table omitted from the DOCX body for layout stability; the full table is available in the Markdown manuscript and stage-2 table files ({len(df)} rows, {len(df.columns)} columns)."
        )
        note.runs[0].italic = True
        note.runs[0].font.size = Pt(9)
        return
    display = df.head(max_rows).copy()
    table = doc.add_table(rows=1, cols=len(display.columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header = table.rows[0].cells
    for idx, col in enumerate(display.columns):
        set_cell_text(header[idx], str(col), bold=True)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "E8EEF7")
        header[idx]._tc.get_or_add_tcPr().append(shading)
    for _, row in display.iterrows():
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], str(value), bold=False)
    doc.add_paragraph()


def add_image(doc: Document, rel_path: str, caption: str) -> None:
    path = MANUSCRIPT / rel_path
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(6.3))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in cap.runs:
        run.font.size = Pt(9)
        run.italic = True


def apply_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Heading 1", 16, RGBColor(31, 78, 121)),
        ("Heading 2", 13, RGBColor(31, 78, 121)),
        ("Heading 3", 11.5, RGBColor(70, 70, 70)),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)


def build_docx(md: str) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    add_page_number(section.footer.paragraphs[0])
    apply_styles(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(TITLE)
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor(31, 78, 121)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(SUBTITLE)
    sub_run.italic = True
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(10.5)
    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    author_run = authors.add_run(AUTHORS)
    author_run.font.name = "Arial"
    author_run.font.size = Pt(10.5)
    doc.add_paragraph()

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("# "):
            i += 1
            continue
        if stripped.startswith("**") and stripped.endswith("**"):
            i += 1
            continue
        if stripped.startswith("Author: "):
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
            i += 1
            continue
        if stripped.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
            if match:
                add_image(doc, match.group(2), match.group(1))
            i += 1
            continue
        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            table_df = parse_markdown_table_lines(table_lines)
            max_rows = 12 if len(table_df) > 12 else len(table_df)
            add_table_to_doc(doc, table_df, max_rows=max_rows)
            continue
        if stripped.startswith("\\[") or stripped.startswith("\\]") or stripped.startswith("\\operatorname"):
            p = doc.add_paragraph(stripped)
            for run in p.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(9)
            i += 1
            continue
        if stripped.startswith("Let a DNA read") or stripped.startswith("For a spaced") or stripped.startswith("The canonical") or stripped.startswith("Thus CSP"):
            p = doc.add_paragraph(stripped)
            i += 1
            continue
        p = doc.add_paragraph(stripped)
        if stripped.startswith(tuple(f"{i}. " for i in range(1, 60))):
            p.paragraph_format.left_indent = Inches(0.15)
            p.paragraph_format.first_line_indent = Inches(-0.15)
        i += 1

    # Use continuous section break to keep final layout stable in Word.
    doc.add_section(WD_SECTION.CONTINUOUS)
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def main() -> None:
    md = build_markdown()
    MD_PATH.write_text(md, encoding="utf-8")
    build_docx(md)
    print(f"Wrote {MD_PATH}")
    print(f"Wrote {DOCX_PATH}")


if __name__ == "__main__":
    main()

