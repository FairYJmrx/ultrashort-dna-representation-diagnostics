from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_MD = PROJECT_ROOT / "docs" / "stage2_literature_analysis_for_advisor_zh.md"
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "stage2_literature_analysis_for_advisor_zh.docx"


TABLE_WIDTH_DXA = 9360


def set_run_font(run, name: str = "Microsoft YaHei", size: float | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for key in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(key), name)


def set_paragraph_spacing(paragraph, before: int = 0, after: int = 80, line: int = 300) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    spacing = ppr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        ppr.append(spacing)
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    spacing.set(qn("w:line"), str(line))
    spacing.set(qn("w:lineRule"), "auto")


def shade_paragraph(paragraph, fill: str = "F4F6F9") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_shading(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW")
        tcpr.append(tcw)
    tcw.set(qn("w:type"), "dxa")
    tcw.set(qn("w:w"), str(width_dxa))


def set_table_width(table, width_dxa: int = TABLE_WIDTH_DXA, indent_dxa: int = 120) -> None:
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "dxa")
    tblw.set(qn("w:w"), str(width_dxa))
    tblind = tblpr.find(qn("w:tblInd"))
    if tblind is None:
        tblind = OxmlElement("w:tblInd")
        tblpr.append(tblind)
    tblind.set(qn("w:type"), "dxa")
    tblind.set(qn("w:w"), str(indent_dxa))


def set_cell_margins(cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    tcmar = tcpr.find(qn("w:tcMar"))
    if tcmar is None:
        tcmar = OxmlElement("w:tcMar")
        tcpr.append(tcmar)
    for margin_name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tcmar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tcmar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_inline_markdown(paragraph, text: str, size: float = 11, bold: bool = False) -> None:
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, "Consolas", size)
            run.font.color.rgb = RGBColor(80, 80, 80)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, "Microsoft YaHei", size)
            run.bold = bold


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color in [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 12, "1F4D78"),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True


def table_widths(col_count: int) -> list[int]:
    presets = {
        2: [4680, 4680],
        3: [2100, 3300, 3960],
        4: [2600, 2900, 1900, 1960],
        5: [1800, 2600, 1900, 1500, 1560],
    }
    if col_count in presets:
        return presets[col_count]
    base = TABLE_WIDTH_DXA // col_count
    widths = [base] * col_count
    widths[-1] += TABLE_WIDTH_DXA - sum(widths)
    return widths


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    norm_rows = [row + [""] * (col_count - len(row)) for row in rows]
    table = doc.add_table(rows=len(norm_rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_table_width(table)
    widths = table_widths(col_count)

    for r_idx, row in enumerate(norm_rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[c_idx])
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, "E8EEF5")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(p, before=0, after=0, line=280)
            add_inline_markdown(p, value.strip(), size=8.2, bold=(r_idx == 0))

    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=80)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip()
        cells = [cell.strip() for cell in raw.strip("|").split("|")]
        is_delimiter = all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)
        if not is_delimiter:
            rows.append(cells)
        i += 1
    return rows, i


def add_paragraph_text(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=80, line=300)
    add_inline_markdown(p, text, size=11)


def add_code_block(doc: Document, code_lines: list[str]) -> None:
    for line in code_lines:
        p = doc.add_paragraph()
        shade_paragraph(p)
        set_paragraph_spacing(p, before=0, after=0, line=280)
        run = p.add_run(line if line else " ")
        set_run_font(run, "Consolas", 9)
    doc.add_paragraph()


def build_docx(md_path: Path, out_path: Path) -> None:
    doc = Document()
    style_document(doc)
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    paragraph_buffer: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_buffer
        if paragraph_buffer:
            add_paragraph_text(doc, " ".join(line.strip() for line in paragraph_buffer))
            paragraph_buffer = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped.startswith("```"):
            flush_paragraph()
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            add_code_block(doc, code_lines)
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            rows, i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            flush_paragraph()
            level = min(len(heading_match.group(1)), 3)
            text = heading_match.group(2).strip()
            if level == 1:
                p = doc.add_paragraph()
                set_paragraph_spacing(p, before=0, after=120, line=320)
                run = p.add_run(text)
                set_run_font(run, "Microsoft YaHei", 18)
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("0B2545")
            else:
                p = doc.add_paragraph(style=f"Heading {level - 1}")
                p.clear()
                add_inline_markdown(p, text, size=16 if level == 2 else 13, bold=True)
            i += 1
            continue

        if stripped.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            set_paragraph_spacing(p, before=0, after=60, line=300)
            add_inline_markdown(p, stripped[2:].strip(), size=11)
            i += 1
            continue

        paragraph_buffer.append(stripped)
        i += 1

    flush_paragraph()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


if __name__ == "__main__":
    build_docx(SOURCE_MD, OUTPUT_DOCX)
    print(OUTPUT_DOCX)
