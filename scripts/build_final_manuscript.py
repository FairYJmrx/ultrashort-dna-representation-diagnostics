from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
MANUSCRIPT = PROJECT_ROOT / "manuscript"
TABLES = MANUSCRIPT / "tables"
FIGS = PROJECT_ROOT / "results" / "figures"
REFERENCES = PROJECT_ROOT / "references" / "references.bib"

TITLE = "Representation diagnostics for ultra-short DNA reads in mNGS-like settings"
SUBTITLE = "A lightweight study of k-mer, canonical spaced-seed, property-aware and attention-compatible encodings"


def read_table_md(path: Path) -> pd.DataFrame:
    text = path.read_text(encoding="utf-8")
    lines = [line for line in text.splitlines() if line.strip().startswith("|")]
    if len(lines) < 2:
        return pd.DataFrame()
    header = [cell.strip() for cell in lines[0].strip("|").split("|")]
    rows = []
    for line in lines[2:]:
        rows.append([cell.strip() for cell in line.strip("|").split("|")])
    return pd.DataFrame(rows, columns=header)


def df_to_markdown(df: pd.DataFrame) -> str:
    if df.empty:
        return ""
    columns = [str(col) for col in df.columns]
    rows = [[str(value) for value in row] for row in df.astype(str).values.tolist()]
    widths = []
    for idx, col in enumerate(columns):
        values = [row[idx] for row in rows]
        widths.append(max([len(col)] + [len(value) for value in values]))
    header = "| " + " | ".join(col.ljust(widths[idx]) for idx, col in enumerate(columns)) + " |"
    sep = "| " + " | ".join("-" * widths[idx] for idx in range(len(columns))) + " |"
    body = ["| " + " | ".join(row[idx].ljust(widths[idx]) for idx in range(len(columns))) + " |" for row in rows]
    return "\n".join([header, sep] + body)


def bib_entries() -> list[dict[str, str]]:
    text = REFERENCES.read_text(encoding="utf-8")
    blocks = re.split(r"\n(?=@)", text.strip())
    entries: list[dict[str, str]] = []
    for block in blocks:
        key_match = re.match(r"@\w+\{([^,]+),", block)
        if not key_match:
            continue
        entry = {"key": key_match.group(1)}
        for field in ["author", "title", "journal", "year", "volume", "pages", "doi", "eprint", "archivePrefix"]:
            match = re.search(rf"\n\s*{field}\s*=\s*\{{(.*?)\}}\s*,?", block, re.IGNORECASE | re.DOTALL)
            if match:
                entry[field] = re.sub(r"\s+", " ", match.group(1)).strip()
        entries.append(entry)
    return entries


def author_year_citation(entries: list[dict[str, str]], key: str) -> str:
    entry = next((item for item in entries if item.get("key") == key), {"key": key})
    authors = entry.get("author", key).split(" and ")
    first = authors[0].split(",")[0]
    suffix = " et al." if len(authors) > 2 else (f" and {authors[1].split(',')[0]}" if len(authors) == 2 else "")
    return f"{first}{suffix}, {entry.get('year', 'n.d.')}"


def ref_text(entry: dict[str, str]) -> str:
    authors = entry.get("author", "").replace(" and ", "; ")
    title = entry.get("title", "")
    year = entry.get("year", "")
    journal = entry.get("journal", "")
    bits = [authors, f"({year})." if year else "", title + "." if title else ""]
    if journal:
        bits.append(journal + ".")
    if entry.get("volume"):
        bits.append(entry["volume"] + ".")
    if entry.get("pages"):
        bits.append(entry["pages"] + ".")
    if entry.get("doi"):
        bits.append("https://doi.org/" + entry["doi"])
    elif entry.get("eprint"):
        prefix = entry.get("archivePrefix", "arXiv")
        bits.append(f"{prefix}:{entry['eprint']}")
    return " ".join(bit for bit in bits if bit).strip()


def build_markdown() -> str:
    refs = bib_entries()
    cite = lambda key: f"({author_year_citation(refs, key)})"

    rep_table = read_table_md(TABLES / "table_representation_taxonomy.md")
    prop_table = read_table_md(TABLES / "table_property_ablation.md")
    param_stable = read_table_md(TABLES / "table_parameter_sensitivity_best_stability.md")
    param_class = read_table_md(TABLES / "table_parameter_sensitivity_best_classification.md")
    attention = read_table_md(TABLES / "table_attention_context_visibility.md")

    prop_short = df_to_markdown(prop_table.head(5))
    param_short = df_to_markdown(param_stable[param_stable["condition"] == "N_3pct"])
    class_short = df_to_markdown(param_class)
    attn_short = df_to_markdown(attention[["length", "observed_layout", "pair_visible_rate", "class_motif_visible_rate", "representation", "macro_f1"]])

    md: list[str] = []
    md.append(f"# {TITLE}\n")
    md.append(f"**{SUBTITLE}**\n")
    md.append("Manuscript draft generated 2026-06-20.\n")
    md.append("## Abstract\n")
    md.append(
        "Ultra-short sequencing reads are common after quality control in clinical metagenomic next-generation sequencing (mNGS), "
        "yet many representation choices for DNA reads are evaluated mainly by downstream classification accuracy. "
        "Here we frame read representation as an information-diagnostic problem: which encodings preserve strand symmetry, local composition, "
        "biochemical properties, perturbation stability, positional context and close-relative separability when reads are only 69-150 bp long? "
        "We compare contiguous k-mers, canonical k-mers, canonical spaced seeds, property-channel encodings, phase-aware encodings and RoPE-like "
        "property encodings on controlled reads and a lightweight close-relative WGS-slice panel from clinically relevant genera. "
        "The main proposed representation, canonical spaced-property encoding (`cspaced_property_l2` in code), concatenates reverse-complement canonical spaced-token counts with low-dimensional DNA property summaries. "
        "It is not a replacement for canonical k-mers. Instead, it acts as a compact, strand-friendly and perturbation-stable auxiliary representation. "
        "In close-relative WGS-slice perturbation audits, adding property summaries to canonical spaced counts improved clean-versus-perturbed cosine by up to 0.028 and reduced L2 perturbation by up to 0.154 at 75 bp under 3% N masking. "
        "By contrast, canonical k-mer and canonical spaced variants remained strong baselines in close-relative classification probes. "
        "A motif-pair diagnostic further showed that read-length effects can be nonlinear for attention-like models: below 150 bp, a class-defining contextual motif pair can be structurally absent, not merely diluted. "
        "The study therefore supports a restrained claim: biologically informed auxiliary encodings can expose robustness and context properties that are hidden by accuracy-only benchmarks, while full mNGS diagnostic claims require larger server-scale validation."
    )

    md.append("## Introduction\n")
    md.append(
        "Clinical mNGS has changed pathogen detection because it can identify unexpected organisms without a fixed target panel "
        f"{cite('Wilson2014')} {cite('Wilson2019')}. However, the computational problem is not just classification. "
        "A clinical read can be short, host-contaminated, quality-trimmed, ambiguous at N positions, or derived from either strand. "
        "These constraints make the representation layer scientifically important: before a classifier can succeed, the encoding must decide what information remains visible."
    )
    md.append(
        "Most mature metagenomic classifiers rely on k-mer or related exact-match signals. Kraken, Kraken 2, CLARK, Centrifuge and Kaiju demonstrate how powerful indexed word or translated-word matching can be at scale "
        f"{cite('Wood2014')} {cite('Wood2019')} {cite('Ounit2015')} {cite('Kim2016')} {cite('Menzel2016')}. "
        "Community benchmarks such as CAMI also show that metagenomic tool performance depends strongly on dataset construction, novelty, abundance and taxonomic difficulty "
        f"{cite('Sczyrba2017')} {cite('Meyer2022')}. This argues against using a small local experiment as a clinical leaderboard."
    )
    md.append(
        "At the same time, DNA language models and self-attention models have made sequence representation a central question. "
        "DeepMicrobes, MetaTransformer, DNABERT, Nucleotide Transformer and HyenaDNA illustrate the move from simple word counts toward learned embeddings, self-attention and single-nucleotide or long-range modeling "
        f"{cite('Liang2020')} {cite('Wichmann2023')} {cite('Ji2021')} {cite('DallaTorre2025')} {cite('Nguyen2023')}. "
        "These models motivate a more precise question for short mNGS reads: which information should be injected before learning, and which information is absent regardless of model capacity?"
    )

    md.append("## Related Work\n")
    md.append(
        "k-mer counting is a foundational alignment-free representation, with efficient counting algorithms and wide use in comparison, classification and sketching "
        f"{cite('Marcais2011')} {cite('Ondov2016')}. Canonical k-mers are not a single named model but a common strand-symmetry operation: a word and its reverse complement are mapped to the same feature index. "
        "Spaced seeds, introduced for sensitive homology search and later adapted to metagenomic classification, provide a mismatch-tolerant alternative to contiguous words "
        f"{cite('Ma2002')} {cite('Brinda2015')}."
    )
    md.append(
        "DNA can also be treated as a signal. Shannon's information theory gives language for capacity, uncertainty and information loss "
        f"{cite('Shannon1948')}, while genomic signal processing and numerical DNA mappings provide precedent for converting bases into biochemical or numeric channels "
        f"{cite('Voss1992')} {cite('Anastassiou2001')} {cite('Cristea2002')}. "
        "Our property encodings follow this tradition: they are deliberately small, interpretable channels rather than learned embeddings."
    )
    md.append(
        "Transformer-style encodings add another issue: position and co-occurrence. Self-attention can connect all observed tokens "
        f"{cite('Vaswani2017')}, and rotary position embeddings provide a compact relative-position mechanism "
        f"{cite('Su2021')}. But attention cannot attend to a motif that has been trimmed away. This distinction motivates our context-visibility diagnostic."
    )

    md.append("## Problem Formulation\n")
    md.append(
        "Let a DNA read be a sequence x = (x1, ..., xL), xi in {A,C,G,T,N}. A representation is a map phi(x) into either a fixed vector or a token sequence. "
        "The paper evaluates phi by information properties rather than by assuming one downstream classifier is definitive."
    )
    md.append(
        "For a contiguous k-mer word w = x_i...x_{i+k-1}, the ordinary count vector stores c_w(x). "
        "The reverse-complement canonical form is canon(w) = min(w, rc(w)) under lexicographic order, so the canonical k-mer count feature is c_canon(w)(x). "
        "This operation is expected to improve strand consistency but may discard strand-specific information."
    )
    md.append(
        "For a spaced seed pattern P = (p1, ..., pm), a spaced token is s_i,P(x) = x_{i+p1}...x_{i+pm}. "
        "A canonical spaced representation counts canon(s_i,P). The proposed canonical spaced-property representation concatenates this count vector with a compact property summary: "
        "mean and standard deviation of hydrogen-bond class, GC indicator, purine indicator and EIIP-like numeric value, plus N fraction, length scaling and sequence entropy. "
        "The final vector is L2-normalized. In code this method is named `cspaced_property_l2`; in the manuscript we call it canonical spaced-property encoding."
    )
    md.append(
        "We evaluate five information properties: compactness, reverse-complement consistency, perturbation stability, read-length/context visibility and close-relative separability. "
        "Accuracy and macro-F1 are used only as tertiary probes: they test whether a simple readout can extract a signal from a representation, not whether the representation is clinically diagnostic."
    )

    md.append("## Experimental Design\n")
    md.append(
        "The local panel contains 21 genomes from six clinically relevant genera: Acinetobacter, Burkholderia, Candida, Enterobacter, Escherichia and Klebsiella. "
        "Reads were sampled as 69, 75, 100, 125 and 150 bp single-end fragments plus a PE150 proxy represented by 300 bp concatenated end information. "
        "Perturbations included reverse complement, 3% N masking and 1% substitution. "
        "The panel is intentionally lightweight and close-relative-biased; it is a stress test, not a universal microbial benchmark."
    )
    md.append(
        "We also used a controlled attention/context diagnostic. Latent templates contain a shared anchor motif near position 18 and a class-specific motif near position 138. "
        "Short reads can include the anchor while excluding the class motif. This design tests whether read shortening removes an entire semantic relation rather than only a proportional number of bases."
    )
    md.append(
        "The metric hierarchy is fixed before interpreting results. Primary metrics are dimensionality, sparsity, paired cosine, perturbation L2 delta, component deltas and motif-pair visibility. "
        "Secondary metrics include close-relative stress probes. Tertiary metrics include accuracy and macro-F1 from nearest-centroid or lightweight linear readouts."
    )

    md.append("## Results\n")
    md.append("### Canonical spaced-property encoding has its clearest advantage in perturbation stability\n")
    md.append(
        "The strongest supported advantage of canonical spaced-property encoding is robustness, especially under N masking. "
        "Adding property summaries to canonical spaced counts consistently increased clean-versus-perturbed cosine and reduced L2 change across 75-300 bp. "
        "The largest effect occurred at 75 bp under 3% N masking: cosine increased by 0.028 and mean L2 perturbation decreased by 0.154. "
        "The effect decreased with read length, which is plausible because longer reads provide more redundant word evidence."
    )
    md.append(prop_short)
    md.append("![Figure 1. Perturbation stability across read lengths.](../results/figures/fig_publication_perturbation_stability.png)")
    md.append("![Figure 2. Component ablation for DNA property summaries.](../results/figures/fig_publication_property_ablation.png)")

    md.append("### Canonicalization, not property summaries alone, explains strand symmetry\n")
    md.append(
        "Reverse-complement robustness was dominated by canonicalization. Canonical contiguous k-mers achieved paired cosine near 1.0 under reverse complement, whereas noncanonical contiguous k-mers had much lower paired cosine. "
        "The canonical spaced-property representation inherits this property from canonical spaced tokens. Therefore the manuscript should not attribute strand invariance to biochemical property summaries alone."
    )

    md.append("### k and spaced-seed pattern sensitivity argues against a single-parameter claim\n")
    md.append(
        "A reviewer would reasonably ask why k=5 or why pattern (0,2,4,6) was selected. We therefore ran a lightweight sensitivity audit over k=4-7 and three spaced patterns. "
        "For N masking, canonical spaced-property variants remained the stability winners across all read lengths in the sampled audit, with paired cosine between 0.994 and 0.997. "
        "However, downstream close-relative probes were parameter-sensitive: the best readout settings changed with task and length. This is evidence for a representation-diagnostics paper, not a universal winner claim."
    )
    md.append(param_short)
    md.append("![Figure 3. Parameter sensitivity for k and spaced-seed patterns.](../results/figures/fig_publication_parameter_sensitivity.png)")

    md.append("### Read length can erase context relations relevant to attention-like models\n")
    md.append(
        "The attention/context diagnostic supports the user's core hypothesis: the loss caused by short reads is not always linear in base count. "
        "At 69, 75 and 100 bp the class-defining motif was absent, pair visibility was 0, and lightweight readouts stayed near chance. "
        "At 125 bp, accidental partial-prefix matches appeared but full pair visibility remained essentially absent. "
        "At 150 bp and PE150, the pair became fully visible and simple readouts reached perfect or near-perfect macro-F1. "
        "A Transformer with self-attention would have many token pairs even at 69 bp, but those pairs cannot include a missing class motif."
    )
    md.append(attn_short)
    md.append("![Figure 4. Context visibility as read length increases.](../results/figures/fig_attention_context_visibility.png)")
    md.append("![Figure 5. Classification probe transition when the motif pair becomes visible.](../results/figures/fig_attention_context_classification.png)")

    md.append("### Close-relative classification is a stress probe, not the paper's main endpoint\n")
    md.append(
        "The close-relative WGS-slice benchmark deliberately tests a harder situation than artificial composition tasks. "
        "It does not support a broad claim that the proposed method is better for species identification. "
        "Canonical k-mer and canonical spaced variants remained strong baselines. In the clean within-genus species probe, the best averaged readout in the core comparison was canonical spaced at 125 bp. "
        "In the target/background probe, canonical spaced-property was best at the PE150 proxy. "
        "In the k/pattern sensitivity audit, best settings varied by task: target/background at 75 bp favored canonical spaced pattern 0-1-3-6, while within-genus probes favored different k-mer or spaced settings depending on length."
    )
    md.append(class_short)
    md.append("![Figure 6. Close-relative WGS-slice classification probes.](../results/figures/fig_publication_close_relative_probes.png)")

    md.append("## Model Suitability\n")
    md.append(
        "Different encodings suggest different model pairings. Canonical k-mer counts remain strong for nearest-centroid, linear and database-index-like methods because they expose exact local composition with strand symmetry. "
        "Canonical spaced counts are useful for compact, mismatch-tolerant linear or centroid readouts. "
        "Canonical spaced-property encoding is best used as an auxiliary dense feature block for QC-like robustness, perturbation-aware screening or concatenation with canonical k-mer features. "
        "Property channels and RoPE-property encodings are more natural for CNNs or attention models because they preserve per-position numeric channels. "
        "The present local study does not prove Transformer superiority; it defines when attention-compatible features have enough observed context to be meaningful."
    )

    md.append("## Discussion\n")
    md.append(
        "The central conclusion is deliberately narrower than a classification claim. Ultra-short DNA read representations differ in what they make stable, compact and visible. "
        "Canonical k-mers remain strong close-relative baselines. Canonical spaced-property encoding contributes a different advantage: compact perturbation stability and interpretability. "
        "This division is scientifically useful because mNGS workflows face both taxonomic discrimination and robustness/QC problems."
    )
    md.append(
        "The study also clarifies the role of accuracy. Accuracy and macro-F1 are helpful only when they are interpreted as readout probes. "
        "If written as clinical endpoints, the current experiments would be underpowered and non-representative. "
        "If written as representation diagnostics, they help map advantage regions and failure modes."
    )
    md.append(
        "Several claims require server-scale follow-up. A larger panel should include many strains per close-relative complex, realistic FASTQ quality profiles, host/background mixtures, abundance variation and repeated random seeds. "
        "A tiny CNN/Transformer comparison should test whether property and RoPE-property channels help when the model can learn local or global interactions. "
        "Kraken2/Centrifuge/Kaiju audits would connect representation diagnostics to clinical-pipeline baselines, and an AMR-gene task would be needed before making resistance-detection claims."
    )

    md.append("## Methods\n")
    md.append(
        "Genome metadata were selected from the local blood-panel spreadsheet and available WGS FASTA sources. Missing close-relative genomes were retrieved using NCBI Datasets when available. "
        "The reproducible scripts generate the close-relative manifest, sampled reads, perturbations, representation matrices, ablation metrics, sensitivity audits, figures and this manuscript. "
        "No heavy neural model was trained locally. Nearest-centroid and lightweight linear probes were used only to test signal readability."
    )
    md.append(
        "Reverse-complement consistency was measured by paired cosine between a clean read and its reverse complement after representation. "
        "Perturbation stability was measured by paired cosine and L2 distance between clean and perturbed representations. "
        "Component ablation compared matched representations that differed by one prior: canonicalization, spaced seeding, property summary, phase or RoPE-like position handling."
    )
    md.append(
        "Parameter sensitivity was run as a sampled audit. Stability used at most 120 paired reads per length/condition, k=4-7 and three spaced patterns. "
        "Classification sensitivity used 75, 150 and PE150 proxy lengths with at most 80 reads per genus. These limits make the audit reproducible on the local computer and should be expanded on a server for stronger statistical inference."
    )

    md.append("## Data and Code Availability\n")
    md.append(
        "The code repository contains the read-generation scripts, representation builders, analysis workflows, result tables, figures and manuscript builder. "
        "Large downloaded reference genomes are not treated as manuscript evidence and should be regenerated from accession manifests when needed. "
        "The repository is intended to archive lightweight reproducibility artifacts while excluding environment folders and bulky downloaded FASTA files."
    )

    md.append("## Limitations\n")
    md.append(
        "The experiments are lightweight and not clinically representative. The close-relative panel has 21 genomes from six genera, so it cannot represent the microbial tree or clinical sample complexity. "
        "The classification probes use simple readouts and are intentionally not clinical performance estimates. "
        "The attention diagnostic is synthetic; it demonstrates structural context loss but does not evaluate a full Transformer. "
        "The property channels are interpretable but may not capture all biologically relevant chemistry or evolutionary constraints."
    )

    md.append("## Conclusions\n")
    md.append(
        "The safest conclusion is that no single DNA representation dominates all short-read mNGS-like settings. "
        "Canonical k-mer remains a strong close-relative classification backbone. "
        "Canonical spaced-property encoding is a compact, strand-friendly and perturbation-stable auxiliary representation with a clear advantage region in robustness diagnostics. "
        "Attention-compatible encodings should be evaluated with explicit context-visibility tests because short reads can delete whole motif relations. "
        "The resulting paper should be positioned as a representation-diagnostics study and a foundation for larger mNGS and AMR validation, not as a final clinical classifier."
    )

    md.append("## References\n")
    for i, entry in enumerate(refs, 1):
        md.append(f"{i}. {ref_text(entry)}")

    return "\n\n".join(md) + "\n"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(table, top=80, start=120, bottom=80, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(TITLE)
    r.bold = True
    r.font.name = "Calibri"
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor.from_string("0B2545")
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(12)
    r2 = p2.add_run(SUBTITLE)
    r2.italic = True
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor.from_string("555555")
    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(18)
    r3 = p3.add_run("Manuscript draft | Representation diagnostics | Generated 2026-06-20")
    r3.font.size = Pt(9)
    r3.font.color.rgb = RGBColor.from_string("666666")


def add_md_table(doc: Document, table_md: str, max_rows: int | None = None) -> None:
    lines = [line for line in table_md.splitlines() if line.strip().startswith("|")]
    if len(lines) < 2:
        return
    header = [cell.strip() for cell in lines[0].strip("|").split("|")]
    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in lines[2:]]
    if max_rows:
        rows = rows[:max_rows]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, item in enumerate(header):
        hdr[i].text = item
        set_cell_shading(hdr[i], "F4F6F9")
        for p in hdr[i].paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(8)
    for row in rows:
        cells = table.add_row().cells
        for i, item in enumerate(row):
            cells[i].text = item
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(8)
    width = 9360 // len(header)
    set_table_width(table, [width] * len(header))
    set_cell_margins(table)
    doc.add_paragraph()


def add_figure(doc: Document, rel: str, caption: str) -> None:
    path = (MANUSCRIPT / rel).resolve() if rel.startswith("..") else (PROJECT_ROOT / rel).resolve()
    if not path.exists():
        para = doc.add_paragraph(f"[Missing figure: {path}]")
        para.style = "Intense Quote"
        return
    doc.add_picture(str(path), width=Inches(6.2))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(10)
    for run in cap.runs:
        run.font.size = Pt(9)
        run.font.italic = True
        run.font.color.rgb = RGBColor.from_string("555555")


def build_docx(md: str, output: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)
    add_title(doc)

    lines = md.splitlines()
    in_refs = False
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            if line[2:].strip() != TITLE:
                doc.add_heading(line[2:].strip(), level=1)
            i += 1
            continue
        if line.startswith("## "):
            title = line[3:].strip()
            in_refs = title == "References"
            doc.add_heading(title, level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            if match:
                add_figure(doc, match.group(2), match.group(1))
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            add_md_table(doc, "\n".join(table_lines), max_rows=12)
            continue
        if re.match(r"^\d+\. ", line) and in_refs:
            p = doc.add_paragraph(style=None)
            p.paragraph_format.left_indent = Inches(0.22)
            p.paragraph_format.first_line_indent = Inches(-0.22)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line)
            run.font.size = Pt(9)
            i += 1
            continue
        if line.startswith("**") and line.endswith("**"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip("*"))
            run.bold = True
            i += 1
            continue
        para = doc.add_paragraph(line)
        para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        i += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    md = build_markdown()
    md_path = MANUSCRIPT / "final_manuscript.md"
    docx_path = MANUSCRIPT / "final_manuscript.docx"
    md_path.write_text(md, encoding="utf-8")
    build_docx(md, docx_path)
    print(f"Wrote {md_path}")
    print(f"Wrote {docx_path}")


if __name__ == "__main__":
    main()
